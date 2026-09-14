from pathlib import Path
from datetime import datetime,timezone
import hashlib
import json
import time
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .scoring import ranked,FACTORS
from .database import audit

TABS = {
 'Ranked Candidates':['Rank','Tie','Application ID','Candidate name','Score','Matched skills','Relevant evidence / projects','Strengths','Requirements not demonstrated','Assessment status','Processed date','CV link'],
 'Score Evidence':['Application ID','Criterion','Weight','Awarded points','Explanation','Cited CV text','Source location'],
 'Needs Review':['Application ID','Filename','Status','Reason','CV link'],
 'Processing Log':['File ID','Application ID','Content version','Rubric version','Model','Job state','Step','Retries','Updated UTC','Error time UTC','Error type','Report revision']}

def safe(value):
    if isinstance(value,str):
        value = ''.join(c for c in value if ord(c)>=32 or c in '\n\t')[:32000]
        if value.lstrip().startswith(('=','+','-','@')) or value.startswith(('\t','\r','\n')):
            return "'"+value
    return value

def utc(value):
    return datetime.fromtimestamp(value,timezone.utc).isoformat(timespec='seconds') if value else ''

def cv_link(file_id):
    return 'https://drive.google.com/file/d/'+file_id+'/view'

def report_rows(db,role_id):
    result = {name:[] for name in TABS}
    role = db.one('SELECT * FROM roles WHERE id=?',(role_id,))
    rubric_row = db.one('SELECT body FROM rubrics WHERE id=?',(role['rubric_id'],))
    criteria = {c['id']:c for c in json.loads(rubric_row['body'])['criteria']} if rubric_row else {}
    for app in ranked(db,role_id):
        findings = json.loads(app['body'])['findings']
        contact = json.loads(app['contact'])
        positive = [f for f in findings if f['level']!='not_demonstrated']
        result['Ranked Candidates'].append([app['rank'],'Tie' if app['tie'] else '',app['id'],contact.get('name') or 'Name not provided',app['score'],
          '; '.join(criteria[f['criterion_id']]['description'] for f in positive),
          '; '.join(c['quote'] for f in positive for c in f['evidence']),
          '; '.join(f['explanation'] for f in findings if f['level']=='supported'),
          '; '.join(criteria[f['criterion_id']]['description']+': '+f['missing_information'] for f in findings if f['level']!='supported'),
          app['status'],utc(app['assessed_at']),cv_link(app['file_id'])])
        for finding in findings:
            criterion = criteria[finding['criterion_id']]
            sources = {s['id']:s['location'] for s in json.loads(app['sections'])}
            for citation in finding['evidence'] or [{'quote':'','section_id':''}]:
                result['Score Evidence'].append([app['id'],criterion['description'],criterion['weight'],
                  criterion['weight']*FACTORS[finding['level']],finding['explanation'],citation['quote'],sources.get(citation['section_id'],'')])
    for app in db.rows("SELECT * FROM applications WHERE role_id=? AND active=1 AND (status!='completed' OR duplicate_of IS NOT NULL)",(role_id,)):
        result['Needs Review'].append([app['id'],app['filename'],app['status'],app['review_reason'] or ('Duplicate of application '+str(app['duplicate_of']) if app['duplicate_of'] else 'Pending assessment'),cv_link(app['file_id'])])
    for job in db.rows('''SELECT j.*,a.file_id FROM jobs j JOIN applications a ON a.id=j.application_id
                          WHERE a.role_id=? ORDER BY j.id''',(role_id,)):
        model = db.one('SELECT model FROM assessments WHERE application_id=? AND version=? AND rubric_id=?',
                       (job['application_id'],job['version'],job['rubric_id']))
        result['Processing Log'].append([job['file_id'],job['application_id'],job['version'],job['rubric_id'],model['model'] if model else '',
          job['state'],job['step'],job['attempts'],utc(job['updated']),utc(job['error_at']),job['error'] or '',role['revision']])
    return result

def build(db,role_id,path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name,rows in report_rows(db,role_id).items():
        sheet = workbook.create_sheet(name)
        sheet.append(TABS[name])
        for row in rows:
            sheet.append([safe(value) for value in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid',fgColor='123B39')
            cell.font = Font(color='FFFFFF',bold=True)
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65,max(16,len(str(column[0].value))+3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top',wrap_text=True)
    workbook.save(path)

def sync(config,db,drive,role_id):
    with db.lease('report:'+role_id) as acquired:
        if not acquired:
            raise RuntimeError('Report synchronization is busy')
        role = db.one('SELECT * FROM roles WHERE id=?',(role_id,))
        revision = role['revision']
        if role['synced_revision']>=revision:
            return
        if not role['report_id']:
            file_id = drive.reserve_id()
            db.execute('UPDATE roles SET report_id=? WHERE id=?',(file_id,role_id))
            role['report_id'] = file_id
        saved = db.one('SELECT * FROM report_revisions WHERE role_id=? AND revision=?',(role_id,revision))
        path = config.data/'reports'/(hashlib.sha256(role_id.encode()).hexdigest()+f'-{revision}.xlsx')
        if not saved or not path.exists():
            # Consistent report read snapshot while other processes continue writing in WAL mode.
            snapshot = Snapshot(db)
            try:
                snap_role = snapshot.one('SELECT revision FROM roles WHERE id=?',(role_id,))
                if snap_role['revision']!=revision:
                    raise RuntimeError('Report revision advanced; retry snapshot')
                build(snapshot,role_id,path)
            finally:
                snapshot.close()
            db.execute('INSERT OR REPLACE INTO report_revisions(role_id,revision,state,path,sha256) VALUES(?,?,?,?,?)',
                       (role_id,revision,'pending',str(path),hashlib.sha256(path.read_bytes()).hexdigest()))
        parent = json.loads(role['folders'])['Reports']
        drive.upload_report(role['report_id'],parent,path,role_id,revision)
        with db.tx() as conn:
            now = time.time()
            conn.execute('UPDATE roles SET synced_revision=MAX(synced_revision,?),uploaded_at=? WHERE id=?',(revision,now,role_id))
            conn.execute("UPDATE report_revisions SET state='uploaded',uploaded_at=? WHERE role_id=? AND revision=?",(now,role_id,revision))
            audit(conn,'report_uploaded',role_id,{'revision':revision,'file_id':role['report_id']})

class Snapshot:
    def __init__(self,db):
        self.conn = db.connect()
        self.conn.execute('BEGIN')
    def rows(self,sql,args=()):
        return [dict(row) for row in self.conn.execute(sql,args)]
    def one(self,sql,args=()):
        rows = self.rows(sql,args)
        return rows[0] if rows else None
    def close(self):
        self.conn.rollback()
        self.conn.close()
