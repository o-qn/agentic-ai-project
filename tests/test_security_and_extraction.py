import json
from pathlib import Path
import pytest
from hr_agent.demo import RUBRIC,drain
from hr_agent.schemas import Assessment
from hr_agent.scoring import validate_evidence,score
from hr_agent.dashboard import create_app
from hr_agent.document_reader import extract,extract_inner,NeedsReview
from hr_agent.database import DB

CV='Name: Example Person\nBuilt a Python application.\nBuilt a SQL database.\n'

def test_assessment_cannot_invent_score_citations_or_coverage():
    sections=[{'id':'s1','text':'Python'}, {'id':'s2','text':'Other evidence'}]
    body={'findings':[{'criterion_id':c['id'],'level':'partial','evidence':[{'section_id':'s1','quote':'Invented Python project'}],
                     'missing_information':'Unclear','explanation':'Some evidence'} for c in RUBRIC['criteria']],
          'covered_sections':['s1']}
    errors=validate_evidence(body,RUBRIC,sections,{'s1'})
    assert any('all extracted' in err for err in errors)
    assert any('citation' in err for err in errors)
    with pytest.raises(ValueError):
        Assessment.model_validate({**body,'score':100})
    with pytest.raises(ValueError):
        score({**body,'findings':body['findings'][:1]},RUBRIC)

def test_untrusted_tool_cannot_change_another_application(system):
    config,db,drive,model,scanner,worker=system
    drive.add('a','a.txt','role-1',CV)
    scanner.run()
    class MaliciousModel:
        def chat(self,*args,**kwargs):
            return {'tool_calls':[{'function':{'name':'read_cv_sections','arguments':{'application_id':9999,'section_ids':['s1']}}},
                                  {'function':{'name':'delete_file','arguments':{'file_id':'another-file'}}}]}
    worker.agent.ollama=MaliciousModel()
    drain(worker)
    assert db.one('SELECT status FROM applications')['status']=='review'
    assert not db.one('SELECT id FROM assessments')
    job=db.one('SELECT agent_state FROM jobs')
    assert json.loads(job['agent_state'])['turns']==12
    assert db.one('SELECT COUNT(*) n FROM rubrics')['n']==3

def test_dashboard_local_boundary_csrf_and_role_scope(system):
    config,db,drive,model,scanner,worker=system
    web=create_app(config,db,model)
    client=web.test_client()
    assert client.get('/').status_code==200
    assert client.post('/api/automatic',json={'enabled':True}).status_code==403
    assert client.get('/',headers={'Host':'attacker.example'}).status_code==400
    assert client.get('/',environ_base={'REMOTE_ADDR':'192.0.2.1'}).status_code==403
    with client.session_transaction() as session:
        csrf=session['csrf']
    assert client.post('/api/automatic',json={'enabled':True},headers={'X-CSRF-Token':csrf}).status_code==200
    assert client.post('/api/automatic',json={'enabled':False},headers={'X-CSRF-Token':csrf,'Origin':'https://attacker.example'}).status_code==403
    assert client.get('/api/roles/no-such-role').status_code==404

def test_audit_is_append_only(system):
    _,db,*_=system
    with pytest.raises(Exception):
        db.execute("UPDATE audit SET event='tampered'")
    with pytest.raises(Exception):
        db.execute('DELETE FROM audit')

def test_scan_lease_prevents_overlap(system):
    _,db,_,_,scanner,_=system
    with db.lease('scanner') as first:
        assert first
        assert scanner.run() is False

def test_txt_and_docx_extraction_with_tables(tmp_path):
    from docx import Document
    txt=tmp_path/'sample.txt';txt.write_text(CV)
    result=extract_inner(txt,40)
    assert result['contact']['name']=='Example Person'
    assert not result['contact']['identity_verified']
    doc=Document();doc.add_paragraph(CV)
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Project';table.cell(0,1).text='SQL ledger'
    path=tmp_path/'sample.docx';doc.save(path)
    assert 'SQL ledger' in extract_inner(path,40)['sections'][0]['text']

def test_long_document_never_silently_truncated(tmp_path):
    text=CV+'Long legitimate work history.\n'*300
    path=tmp_path/'long.txt';path.write_text(text)
    result=extract_inner(path,40)
    assert ''.join(s['text'] for s in result['sections'])==text
    path.write_text('X'*300000)
    with pytest.raises(NeedsReview,match='budget'):
        extract_inner(path,40)

def test_blank_pdf_and_corrupt_docx_require_review(system,tmp_path):
    from pypdf import PdfWriter
    config,*_=system
    writer=PdfWriter();writer.add_blank_page(width=300,height=400)
    path=tmp_path/'blank.pdf'
    with path.open('wb') as file:
        writer.write(file)
    with pytest.raises(NeedsReview):
        extract(path,config)
    path=tmp_path/'corrupt.docx';path.write_bytes(b'not a docx')
    with pytest.raises(NeedsReview):
        extract(path,config)

def test_restart_after_extraction_skips_download(system):
    config,db,drive,model,scanner,worker=system
    drive.add('a','a.txt','role-1',CV);scanner.run()
    worker.tick();worker.tick()
    drive.fail['download']=10
    from hr_agent.job_queue import Worker
    restarted=Worker(config,DB(config.data),drive,model)
    drain(restarted)
    assert db.one('SELECT status FROM applications')['status']=='completed'
    assert drive.fail['download']==10
